#!/usr/bin/env python3
"""
Xuất CSV từ file Excel dự án EMR (nhiều sheet, header không đồng nhất) sang
đúng định dạng cột mà trang /emr/nhap-lieu (ImportCsvPanel) mong đợi.

CHẠY Ở MÁY CỦA BẠN, KHÔNG chạy trên máy chủ dùng chung — script này đọc file
Excel có thể chứa dữ liệu nội bộ/nhạy cảm và ghi CSV ra thư mục exports/ (đã
có trong .gitignore của module này). KHÔNG commit các file CSV kết quả lên Git.

Cách dùng:
    pip install openpyxl
    python scripts/export_excel_to_csv.py "/duong/dan/file_du_an_EMR.xlsx"

Sẽ tạo (tùy sheet có tồn tại trong file):
    exports/quy_trinh.csv              -> nhập vào "Quy trình - tài liệu"
    exports/thiet_bi_yte.csv           -> nhập vào "Thiết bị y tế"
    exports/chu_ky_so.csv              -> nhập vào "Chữ ký số" (KHÔNG xuất số CCCD)
    exports/thiet_bi_cntt.csv          -> nhập vào "Thiết bị CNTT"
    exports/bieu_mau_x_khoa.csv        -> nhập vào "Ma trận biểu mẫu x khoa"

Tên sheet có thể khác nhau giữa các phiên bản file — chỉnh SHEET_NAMES bên dưới
nếu script không tìm thấy sheet tương ứng.
"""

import csv
import sys
from pathlib import Path

try:
    import openpyxl
except ImportError:
    sys.exit("Cần cài openpyxl trước: pip install openpyxl")

SHEET_NAMES = {
    "quy_trinh": "Ds Quy trình",
    "tbyt": "Ds TBYT",
    "cks": "Ds CKS",
    "cntt": "Ds CNTT",
    "matrix": "Ds Biểu mẫu (Vinh)",
}

# Khớp tên khoa/phòng viết hoa có đánh số trong file gốc (vd: "1. KHOA KHÁM BỆNH")
# sang đúng tên đã seed trong bảng emr_departments (xem migration SQL).
DEPARTMENT_NAME_MAP = {
    "KHOA KHÁM BỆNH": "Khoa Khám bệnh",
    "KHOA HỒI SỨC": "Khoa Hồi sức",
    "KHOA CẤP CỨU": "Khoa Cấp cứu",
    "KHOA NỘI": "Khoa Nội",
    "KHOA NGOẠI": "Khoa Ngoại",
    "KHOA SẢN": "Khoa Sản",
    "KHOA NHI": "Khoa Nhi",
    "KHOA PT-GMHS": "Khoa Phẫu thuật - GMHS",
    "KHOA VIP": "Khoa VIP",
    "ĐV THẬN NHÂN TẠO": "Đơn vị Thận nhân tạo",
    "ĐV NỘI SOI TIÊU HÓA": "Đơn vị Nội soi tiêu hóa",
    "ĐV VLTL - PHCN": "Đơn vị VLTL - PHCN",
    "ĐV RHM": "Đơn vị Răng Hàm Mặt",
}


def strip_number_prefix(label: str) -> str:
    """"1. KHOA KHÁM BỆNH" -> "KHOA KHÁM BỆNH" """
    parts = label.split(".", 1)
    if len(parts) == 2 and parts[0].strip().isdigit():
        return parts[1].strip()
    return label.strip()


def write_csv(path: Path, header: list[str], rows: list[list[str]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(header)
        w.writerows(rows)
    print(f"  -> {path}  ({len(rows)} dòng)")


def col_value(ws, row_idx: int, header_row: list, name: str):
    if name not in header_row:
        return None
    col_idx = header_row.index(name) + 1
    v = ws.cell(row=row_idx, column=col_idx).value
    return v


def export_quy_trinh(wb, out_dir: Path):
    name = SHEET_NAMES["quy_trinh"]
    if name not in wb.sheetnames:
        return
    ws = wb[name]
    header_row = [c.value for c in ws[1]]
    rows = []
    for r in range(2, ws.max_row + 1):
        doc_name = col_value(ws, r, header_row, "Tên tài liệu Yêu cầu")
        if not doc_name:
            continue
        dept = col_value(ws, r, header_row, "Khoa/ phòng (dự kiến)") or ""
        status = col_value(ws, r, header_row, "Trạng thái") or "Dự thảo"
        deadline = col_value(ws, r, header_row, "Deadline")
        notes = col_value(ws, r, header_row, "Ghi chú") or ""
        rows.append([str(doc_name), str(dept), str(status), str(deadline or ""), str(notes)])
    write_csv(out_dir / "quy_trinh.csv", ["name", "department", "status", "deadline", "notes"], rows)


def export_tbyt(wb, out_dir: Path):
    name = SHEET_NAMES["tbyt"]
    if name not in wb.sheetnames:
        return
    ws = wb[name]
    header_row = [c.value for c in ws[1]]
    rows = []
    for r in range(2, ws.max_row + 1):
        dev_name = col_value(ws, r, header_row, "Tên thiết Bị YTế")
        if not dev_name:
            continue
        model = col_value(ws, r, header_row, "Model") or ""
        manufacturer = col_value(ws, r, header_row, "Hãng SX") or ""
        dept = col_value(ws, r, header_row, "Khoa") or ""
        img_fmt = col_value(ws, r, header_row, "Đuôi Hình lưu Trữ") or ""
        integration = col_value(ws, r, header_row, "Phương Án Bắt Hình") or ""
        status_raw = str(col_value(ws, r, header_row, "Trạng thái") or "").strip()
        status = "Đã thực hiện" if "đã" in status_raw.lower() else (
            "Đang thực hiện" if status_raw else "Chưa thực hiện"
        )
        rows.append([str(dev_name), str(model), str(manufacturer), str(dept), str(img_fmt), str(integration), status])
    write_csv(
        out_dir / "thiet_bi_yte.csv",
        ["name", "model", "manufacturer", "department", "image_format", "integration_method", "status"],
        rows,
    )


def export_cks(wb, out_dir: Path):
    """Xuất danh sách chữ ký số — CỐ Ý bỏ cột CCCD, không xuất ra CSV."""
    name = SHEET_NAMES["cks"]
    if name not in wb.sheetnames:
        return
    ws = wb[name]
    # Sheet này có dòng tiêu đề phụ ở row 1, header thật ở row 3 theo cấu trúc quan sát được.
    header_row = [c.value for c in ws[3]]
    if "Họ và Tên" not in header_row:
        header_row = [c.value for c in ws[1]]
    rows = []
    for r in range(4, ws.max_row + 1):
        staff_name = col_value(ws, r, header_row, "Họ và Tên")
        if not staff_name:
            continue
        dept = col_value(ws, r, header_row, "Khoa Phòng") or ""
        title = col_value(ws, r, header_row, "Chức Danh") or ""
        provider = col_value(ws, r, header_row, "Đơn vị \ncung cấp") or col_value(ws, r, header_row, "Đơn vị cung cấp") or ""
        rows.append([str(staff_name), str(dept), str(title), str(provider), "", ""])
        # Ngày cấp/hết hạn nằm ở các cột không tên cố định tùy phiên bản file —
        # điền tay sau khi nhập, hoặc chỉnh script này theo đúng cột thật của bạn.
    write_csv(
        out_dir / "chu_ky_so.csv",
        ["staff_name", "department", "title", "provider", "issued_at", "expires_at"],
        rows,
    )
    print("  LƯU Ý: đã cố ý KHÔNG xuất số CCCD/CMND ra file này.")


def export_cntt(wb, out_dir: Path):
    name = SHEET_NAMES["cntt"]
    if name not in wb.sheetnames:
        return
    ws = wb[name]
    locations = [c.value for c in ws[1]][1:]  # bỏ cột đầu "Vị trí"
    kinds = [c.value for c in ws[2]][1:]      # "Hiện có" / "Bổ sung", song song với locations
    rows_out = []
    for r in range(3, ws.max_row + 1):
        device = ws.cell(row=r, column=1).value
        if not device:
            continue
        # gộp theo location: mỗi location có 2 cột liền nhau (Hiện có, Bổ sung)
        seen = {}
        for i, (loc, kind) in enumerate(zip(locations, kinds)):
            if not loc:
                continue
            val = ws.cell(row=r, column=i + 2).value or 0
            seen.setdefault(loc, {"available": 0, "needed": 0})
            if kind == "Hiện có":
                seen[loc]["available"] = val
            elif kind == "Bổ sung":
                seen[loc]["needed"] = val
        for loc, v in seen.items():
            if v["available"] or v["needed"]:
                rows_out.append([str(device), str(loc), str(int(v["available"] or 0)), str(int(v["needed"] or 0)), ""])
    write_csv(out_dir / "thiet_bi_cntt.csv", ["device_name", "location", "qty_available", "qty_needed", "notes"], rows_out)


def export_matrix(wb, out_dir: Path):
    name = SHEET_NAMES["matrix"]
    if name not in wb.sheetnames:
        return
    ws = wb[name]
    header_row = [c.value for c in ws[1]]

    dept_cols = []  # (col_idx_progress, canonical_department_name)
    for i, h in enumerate(header_row):
        if h == "Tiến độ triển khai" and i > 0:
            raw_dept_label = header_row[i - 1]
            if not raw_dept_label:
                continue
            plain = strip_number_prefix(str(raw_dept_label)).upper()
            canonical = DEPARTMENT_NAME_MAP.get(plain)
            if canonical:
                dept_cols.append((i + 1, canonical))  # openpyxl cột 1-based

    if not dept_cols:
        print("  Không nhận diện được cột khoa/phòng nào trong sheet ma trận — bỏ qua.")
        return

    dept_names = [d[1] for d in dept_cols]
    rows_out = []
    for r in range(2, ws.max_row + 1):
        form_name = col_value(ws, r, header_row, "Tên biểu mẫu")
        if not form_name:
            continue
        row_vals = [str(form_name)]
        for col_idx, _dept in dept_cols:
            val = ws.cell(row=r, column=col_idx).value
            row_vals.append(str(val) if val else "")
        if any(row_vals[1:]):
            rows_out.append(row_vals)

    write_csv(out_dir / "bieu_mau_x_khoa.csv", ["Tên biểu mẫu", *dept_names], rows_out)


def main():
    if len(sys.argv) != 2:
        sys.exit("Dùng: python scripts/export_excel_to_csv.py <đường dẫn file Excel dự án>")
    src = Path(sys.argv[1])
    if not src.exists():
        sys.exit(f"Không tìm thấy file: {src}")

    out_dir = Path(__file__).resolve().parent.parent / "exports"
    print(f"Đọc {src} ...")
    wb = openpyxl.load_workbook(src, data_only=True)

    export_quy_trinh(wb, out_dir)
    export_tbyt(wb, out_dir)
    export_cks(wb, out_dir)
    export_cntt(wb, out_dir)
    export_matrix(wb, out_dir)

    print(f"\nXong. Các file CSV nằm trong: {out_dir}")
    print("Nhập từng file qua trang /emr/nhap-lieu. KHÔNG commit thư mục exports/ lên Git.")


if __name__ == "__main__":
    main()
